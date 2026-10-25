"""Аналитика и отчёты для руководства.

GET /analytics?group=complex — по ЖК (или building, judge, court, legal_entity, responsible, year, subject):
сколько дел, заявлено, присуждено, % присуждённого, доля дел с обжалованием, конфликты, открытые и
просроченные сроки. Все фильтры реестра (CaseFilter) работают и здесь.
GET /analytics/load — нагрузка юристов. XLSX — /reports/analytics.xlsx; еженедельный отчёт — по расписанию
([documents] report_weekday / report_hour), список — /reports.
"""
import io
from collections import defaultdict
from datetime import timedelta
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse, Response
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .. import clock
from ..config import get_settings
from ..db import get_db
from ..filters import CaseFilter, case_year
from ..logic import log
from ..models import Assignment, Case, CaseValue, Court, Deadline, Enforcement, Report, User
from ..notify import notify
from ..security import Principal, require
from .cases import DISPUTED, awarded_sq, claimed_sq, conflicts_sq

router = APIRouter(prefix="/api/v1", tags=["Аналитика и отчёты"])

GROUPS = {"complex": "ЖК", "building": "Корпус", "judge": "Судья", "court": "Суд", "legal_entity": "Юрлицо",
          "responsible": "Ответственный", "year": "Год дела", "subject": "Предмет иска"}


def _judges(db: Session) -> dict[int, str]:
    """Судья по делу — итоговое значение поля judge (спорные — «не определён»)."""
    rows = db.execute(select(CaseValue.case_id, CaseValue.value, CaseValue.status).where(CaseValue.field == "judge")).all()
    return {cid: (v if st not in DISPUTED and v else "не определён") for cid, v, st in rows}


def analytics_rows(db: Session, f: CaseFilter, group: str) -> list[dict]:
    if group not in GROUPS:
        raise HTTPException(400, f"group: {', '.join(GROUPS)}")
    today = clock.today()
    open_sq = (select(func.count()).where(Deadline.case_id == Case.id, Deadline.status == "open")
               .correlate(Case).scalar_subquery())
    over_sq = (select(func.count()).where(Deadline.case_id == Case.id, Deadline.status == "open", Deadline.due_date < today)
               .correlate(Case).scalar_subquery())
    rows = db.execute(select(Case.id, Case.complex, Case.building, Case.court_code, Case.court_name, Case.legal_entity,
                             Case.responsible, Case.subject, case_year, Case.appeal_flag, awarded_sq, claimed_sq,
                             conflicts_sq, open_sq, over_sq).where(*f.where())).all()
    judges = _judges(db) if group == "judge" else {}
    courts = {c.code: c.short_name or c.name for c in db.scalars(select(Court))} if group == "court" else {}
    acc: dict[str, dict] = defaultdict(lambda: {"cases": 0, "with_award": 0, "claimed": 0.0, "awarded": 0.0, "claimed_n": 0,
                                                "appeal": 0, "conflicts": 0, "open_deadlines": 0, "overdue": 0})
    for r in rows:
        (cid, cx, b, cc, cn, le, resp, subj, yr, appeal, aw, cl, conf, od, ov) = r
        key = {"complex": cx, "building": b, "court": courts.get(cc) or cn or cc, "legal_entity": le, "responsible": resp,
               "subject": subj, "year": yr, "judge": judges.get(cid, "не определён")}[group] or "не указано"
        a = acc[key]
        a["cases"] += 1
        a["appeal"] += bool(appeal)
        a["conflicts"] += conf or 0
        a["open_deadlines"] += od or 0
        a["overdue"] += ov or 0
        if aw is not None:
            a["with_award"] += 1
            a["awarded"] += float(aw)
        if cl is not None and aw is not None:  # % присуждённого — только по делам, где известны обе суммы
            a["claimed"] += float(cl)
            a["claimed_n"] += 1
    out = []
    for key, a in acc.items():
        aw_with_claim = a["awarded"] if a["claimed_n"] == a["with_award"] else None
        pct = round(100 * aw_with_claim / a["claimed"], 1) if aw_with_claim is not None and a["claimed"] else None
        out.append({"key": key, **a, "awarded": round(a["awarded"], 2), "claimed": round(a["claimed"], 2),
                    "awarded_pct": pct, "appeal_pct": round(100 * a["appeal"] / a["cases"], 1)})
    return sorted(out, key=lambda x: (-x["cases"], x["key"]))


@router.get("/analytics")
def analytics(group: str = "complex", f: CaseFilter = Depends(CaseFilter.dep), db: Session = Depends(get_db),
              _: Principal = Depends(require(role="viewer", scope="read"))):
    rows = analytics_rows(db, f, group)
    total = {k: sum(r[k] for r in rows) for k in ("cases", "with_award", "awarded", "claimed", "appeal", "conflicts",
                                                  "open_deadlines", "overdue")}
    return {"group": group, "group_label": GROUPS[group], "rows": rows, "total": total, "groups": GROUPS}


def lawyer_load(db: Session) -> list[dict]:
    today = clock.today()
    users = {u.login: u.full_name or u.login for u in db.scalars(select(User).where(User.active, User.role != "viewer"))}
    cases = dict(db.execute(select(Case.responsible, func.count()).where(Case.responsible != "").group_by(Case.responsible)).all())
    dls = dict(db.execute(select(Case.responsible, func.count()).join(Deadline, Deadline.case_id == Case.id)
                          .where(Deadline.status == "open", Case.responsible != "").group_by(Case.responsible)).all())
    over = dict(db.execute(select(Case.responsible, func.count()).join(Deadline, Deadline.case_id == Case.id)
                           .where(Deadline.status == "open", Deadline.due_date < today, Case.responsible != "")
                           .group_by(Case.responsible)).all())
    week = dict(db.execute(select(Case.responsible, func.count()).join(Deadline, Deadline.case_id == Case.id)
                           .where(Deadline.status == "open", Deadline.due_date >= today,
                                  Deadline.due_date <= today + timedelta(days=7), Case.responsible != "")
                           .group_by(Case.responsible)).all())
    conf = dict(db.execute(select(Case.responsible, func.count()).join(CaseValue, CaseValue.case_id == Case.id)
                           .where(CaseValue.status.in_(DISPUTED), Case.responsible != "").group_by(Case.responsible)).all())
    asg = dict(db.execute(select(Assignment.assignee, func.count()).where(Assignment.status == "open")
                          .group_by(Assignment.assignee)).all())
    logins = sorted(set(users) | set(cases) | set(asg))
    return [{"login": lg, "name": users.get(lg, lg), "cases": cases.get(lg, 0), "open_deadlines": dls.get(lg, 0),
             "week": week.get(lg, 0), "overdue": over.get(lg, 0), "conflicts": conf.get(lg, 0), "assignments": asg.get(lg, 0)}
            for lg in logins]


@router.get("/analytics/load")
def load(db: Session = Depends(get_db), _: Principal = Depends(require(role="viewer", scope="read"))):
    return lawyer_load(db)


# --- XLSX ------------------------------------------------------------------------------------

def build_xlsx(db: Session, f: CaseFilter | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    f = f or CaseFilter()
    wb = Workbook()
    first = True
    for group in ("complex", "judge", "responsible", "legal_entity"):
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = f"По {GROUPS[group]}"[:31]  # двоеточие в имени листа Excel недопустимо
        ws.append([GROUPS[group], "Дел", "С присуждённым", "Заявлено, ₽", "Присуждено, ₽", "% присуждено", "Обжалование, %",
                   "Конфликты", "Открытые сроки", "Просрочено"])
        for r in analytics_rows(db, f, group):
            ws.append([r["key"], r["cases"], r["with_award"], r["claimed"], r["awarded"], r["awarded_pct"], r["appeal_pct"],
                       r["conflicts"], r["open_deadlines"], r["overdue"]])
        _style(ws, Font)
    ws = wb.create_sheet("Нагрузка юристов")
    ws.append(["Юрист", "Логин", "Дел", "Открытые сроки", "На неделе", "Просрочено", "Конфликты", "Поручения"])
    for r in lawyer_load(db):
        ws.append([r["name"], r["login"], r["cases"], r["open_deadlines"], r["week"], r["overdue"], r["conflicts"], r["assignments"]])
    _style(ws, Font)
    ws = wb.create_sheet("Исполнение")
    ws.append(["Дело", "Сторона", "Лист", "ИП", "Статус", "Сумма, ₽", "Взыскано, ₽", "Срок предъявления"])
    for e in db.scalars(select(Enforcement).order_by(Enforcement.id)):
        c = db.get(Case, e.case_id)
        ws.append([c.case_number if c else e.case_id, "с нас" if e.side == "debtor" else "в нашу пользу", e.writ_number,
                   e.ip_number, e.status, float(e.amount or 0), float(e.recovered or 0),
                   e.presentation_until.strftime("%d.%m.%Y") if e.presentation_until else ""])
    _style(ws, Font)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _style(ws, Font) -> None:
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(48, max(len(str(c.value or "")) for c in col) + 2))


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@router.get("/reports/analytics.xlsx")
def analytics_xlsx(f: CaseFilter = Depends(CaseFilter.dep), db: Session = Depends(get_db),
                   who: Principal = Depends(require(role="viewer", scope="export"))):
    data = build_xlsx(db, f)
    log(db, who, "report_downloaded", "report", None, kind="analytics")
    db.commit()
    name = f"Аналитика_{clock.today().isoformat()}.xlsx"
    return Response(data, media_type=XLSX, headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})


# --- Отчёты по расписанию ---------------------------------------------------------------------

def reports_dir() -> Path:
    return Path(get_settings().reports_dir)


def weekly_report(db: Session, force: bool = False) -> Report | None:
    """Еженедельный отчёт: в [documents] report_weekday после report_hour — один раз за неделю."""
    s = get_settings()
    now = clock.local(clock.now())
    y, w, wd = now.isocalendar()
    key = f"{y}-W{w:02d}"
    if not force and (wd != s.report_weekday or now.hour < s.report_hour):
        return None
    if not force and db.scalar(select(Report.id).where(Report.kind == "weekly", Report.period_key == key)):
        return None
    d = reports_dir()
    d.mkdir(parents=True, exist_ok=True)
    name = f"Еженедельный_отчёт_{now.date().isoformat()}{'_' + now.strftime('%H%M') if force else ''}.xlsx"
    (d / name).write_bytes(build_xlsx(db))
    r = Report(kind="weekly" if not force else "manual", title=f"Отчёт за неделю {key}", file=name, period_key=key)
    db.add(r)
    db.flush()
    admins = db.scalars(select(User.login).where(User.active, User.role == "admin")).all()
    notify(db, admins, "report", f"Готов еженедельный отчёт ({now.strftime('%d.%m.%Y')})", link="/analytics?tab=reports",
           dedup=f"report:{r.id}")
    return r


@router.get("/reports")
def list_reports(db: Session = Depends(get_db), _: Principal = Depends(require(role="viewer"))):
    return [{"id": r.id, "kind": r.kind, "title": r.title, "file": r.file, "created_at": r.created_at.isoformat()}
            for r in db.scalars(select(Report).order_by(Report.id.desc()).limit(100))]


@router.post("/reports/weekly")
def make_weekly(db: Session = Depends(get_db), who: Principal = Depends(require(role="admin"))):
    """Сформировать отчёт сейчас (не дожидаясь расписания)."""
    r = weekly_report(db, force=True)
    log(db, who, "report_created", "report", r.id, file=r.file)
    db.commit()
    return {"id": r.id, "file": r.file}


@router.get("/reports/{report_id}/file")
def report_file(report_id: int, db: Session = Depends(get_db), _: Principal = Depends(require(role="viewer"))):
    r = db.get(Report, report_id)
    if r is None or not (reports_dir() / r.file).is_file():
        raise HTTPException(404, "Файл отчёта не найден")
    return FileResponse(reports_dir() / r.file, media_type=XLSX, filename=r.file)
