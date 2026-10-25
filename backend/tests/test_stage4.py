"""Этап 4: исполнительные производства, аналитика, отчёты."""
from datetime import date

from test_api import ingest_case, obs


def test_enforcement_deadlines(svc, lawyer):
    cid = ingest_case(svc)
    r = lawyer.post("/api/v1/enforcement", json={"case_id": cid, "side": "debtor", "ip_number": "12345/26/50001-ИП",
                                                  "status": "возбуждено", "received_date": "2027-01-11", "amount": "141611.96"})
    assert r.status_code == 201, r.text
    eid = r.json()["id"]
    dls = [d for d in lawyer.get(f"/api/v1/cases/{cid}").json()["deadlines"] if d["kind"] == "voluntary"]
    assert len(dls) == 1 and dls[0]["due_date"] == "2027-01-18"  # 5 рабочих дней от 11.01.2027 (пн)
    lawyer.patch(f"/api/v1/enforcement/{eid}", json={"status": "окончено", "recovered": "141611.96"})
    dls = [d for d in lawyer.get(f"/api/v1/cases/{cid}").json()["deadlines"] if d["kind"] == "voluntary"]
    assert dls[0]["status"] == "cancelled"
    # взыскатель: срок предъявления — 3 года от вступления в силу
    r = lawyer.post("/api/v1/enforcement", json={"case_id": cid, "side": "creditor", "writ_number": "ФС 0123",
                                                  "effective_date": "2027-02-01"}).json()
    assert r["presentation_until"] == "2030-02-01"
    assert any(d["kind"] == "writ" for d in lawyer.get(f"/api/v1/cases/{cid}").json()["deadlines"])
    assert lawyer.post("/api/v1/enforcement", json={"case_id": cid, "status": "ерунда"}).status_code == 400
    # загрузка из ФССП: тот же номер ИП — обновление
    res = svc.post("/api/v1/ingest/enforcement", json=[{"uid": "50RS0010-01-2025-000287-27", "ip_number": "12345/26/50001-ИП",
                                                        "status": "окончено", "recovered": "100"}]).json()
    assert res[0]["ok"] and res[0]["created"] is False


def test_analytics_and_load(svc, lawyer):
    ingest_case(svc, complex="ЖК Героев", responsible="lawyer1")
    obs(svc, "claim_total", "1С", "1000000")
    obs(svc, "amount_debt", "1С", "400000")
    obs(svc, "amount_penalty", "1С", "100000")
    svc.post("/api/v1/ingest/cases", json=[{"uid": "U-9", "case_number": "2-9/27", "complex": "ЖК Героев", "appeal_flag": True}])
    a = lawyer.get("/api/v1/analytics", params={"group": "complex"}).json()
    row = a["rows"][0]
    assert row["key"] == "ЖК Героев" and row["cases"] == 2 and row["awarded"] == 500000 and row["awarded_pct"] == 50.0
    assert row["appeal_pct"] == 50.0
    assert lawyer.get("/api/v1/analytics", params={"group": "nope"}).status_code == 400
    assert lawyer.get("/api/v1/analytics", params={"group": "judge"}).status_code == 200
    load = {x["login"]: x for x in lawyer.get("/api/v1/analytics/load").json()}
    assert load["lawyer1"]["cases"] == 1


def test_reports(svc, lawyer, admin, tmp_path, monkeypatch):
    from openpyxl import load_workbook
    import io
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "reports_dir", str(tmp_path))
    ingest_case(svc)
    r = lawyer.get("/api/v1/reports/analytics.xlsx")
    assert r.status_code == 200 and "Нагрузка юристов" in load_workbook(io.BytesIO(r.content)).sheetnames
    assert lawyer.post("/api/v1/reports/weekly").status_code == 403
    rid = admin.post("/api/v1/reports/weekly").json()["id"]
    assert admin.get(f"/api/v1/reports/{rid}/file").status_code == 200
    assert any("отчёт" in n["title"] for n in admin.get("/api/v1/notifications").json()["items"])
    _ = date
