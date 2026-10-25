"""Этап 3: печатные формы, поиск по текстам документов."""
import io
import json

from docx import Document

from test_api import ingest_case, obs


def upload(svc, text, name, case=None):
    meta = {"filename": name, "text": text, "act_type": "решение", **({"case": case} if case else {})}
    r = svc.post("/api/v1/ingest/documents/upload", files={"file": (name, b"%PDF-1.4 " + name.encode())},
                 data={"meta": json.dumps(meta)})
    assert r.status_code == 200, r.text
    return r.json()["document_id"]


def test_fulltext_search_russian_morphology(svc, viewer):
    cid = ingest_case(svc)
    upload(svc, "Суд решил взыскать неустойку за просрочку передачи квартиры.", "a.pdf", {"case_id": cid})
    upload(svc, "Определение о судебных расходах.", "b.pdf")
    r = viewer.get("/api/v1/documents/search", params={"q": "неустойка просрочка"}).json()
    assert len(r) == 1 and r[0]["case_id"] == cid and "\x02" in r[0]["snippet"]
    assert viewer.get("/api/v1/documents/search", params={"q": "\"судебных расходах\""}).json()[0]["filename"] == "b.pdf"
    assert viewer.get("/api/v1/documents/search", params={"q": "аренда"}).json() == []


def test_forms_list_generate_upload(svc, lawyer, admin):
    cid = ingest_case(svc, plaintiffs="Орлов А.А.")
    obs(svc, "amount_debt", "1С", "141611.96")
    forms = lawyer.get("/api/v1/forms").json()
    assert {f["key"] for f in forms} >= {"Ходатайство_об_отложении", "Заявление_о_выдаче_копии_решения"}
    r = lawyer.post(f"/api/v1/cases/{cid}/forms/Заявление_о_выдаче_копии_решения")
    assert r.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "2-1459/2025" in text and "Орлов А.А." in text and "141 611 руб. 96 коп." in text and "{{" not in text
    assert "{{ v.amount_debt }}" in lawyer.get("/api/v1/forms/variables").json()
    # свой шаблон администратора
    d = Document()
    d.add_paragraph("Дело {{ case.case_number }}, истец {{ case.plaintiffs }}")
    buf = io.BytesIO()
    d.save(buf)
    assert lawyer.post("/api/v1/forms", files={"file": ("Мой.docx", buf.getvalue())}).status_code == 403
    assert admin.post("/api/v1/forms", files={"file": ("Тест_формы.docx", buf.getvalue())}).status_code == 201
    out = lawyer.post(f"/api/v1/cases/{cid}/forms/Тест_формы")
    assert "Дело 2-1459/2025, истец Орлов А.А." in Document(io.BytesIO(out.content)).paragraphs[0].text
    assert admin.delete("/api/v1/forms/Тест_формы").status_code == 204
    assert admin.post("/api/v1/forms", files={"file": ("x.txt", b"hello")}).status_code == 400
